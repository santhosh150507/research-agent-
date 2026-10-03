from sqlalchemy.orm import Session
from sqlalchemy import select
from app.models.literature_review import LiteratureReview
from app.core.errors import AppException
from fastapi.responses import Response
import io

def export_review(db: Session, user_id: int, review_id: int, format: str):
    review = db.execute(select(LiteratureReview).where(LiteratureReview.id == review_id)).scalar_one_or_none()
    if not review:
        raise AppException("Review not found", status_code=404)
        
    title = review.title or "Literature Review"
    # Structure of review content is roughly sections and references
    content = review.content or {}
    sections = content.get("sections", [])
    
    if format == "md":
        md = f"# {title}\n\n"
        for s in sections:
            md += f"## {s.get('heading', '')}\n"
            for b in s.get('blocks', []):
                md += f"{b.get('text', '')}\n\n"
        
        return Response(content=md, media_type="text/markdown", headers={"Content-Disposition": f'attachment; filename="{title}.md"'})
        
    elif format == "docx":
        from docx import Document
        doc = Document()
        doc.add_heading(title, 0)
        for s in sections:
            doc.add_heading(s.get('heading', ''), level=1)
            for b in s.get('blocks', []):
                doc.add_paragraph(b.get('text', ''))
                
        f = io.BytesIO()
        doc.save(f)
        f.seek(0)
        return Response(content=f.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="{title}.docx"'})
        
    elif format == "pdf":
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        
        f = io.BytesIO()
        doc = SimpleDocTemplate(f, pagesize=letter)
        styles = getSampleStyleSheet()
        Story = []
        
        Story.append(Paragraph(title, styles["Title"]))
        Story.append(Spacer(1, 12))
        
        for s in sections:
            Story.append(Paragraph(s.get('heading', ''), styles["Heading1"]))
            for b in s.get('blocks', []):
                Story.append(Paragraph(b.get('text', ''), styles["Normal"]))
                Story.append(Spacer(1, 12))
                
        doc.build(Story)
        f.seek(0)
        return Response(content=f.getvalue(), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{title}.pdf"'})
        
    else:
        raise AppException("Invalid format", status_code=400)
